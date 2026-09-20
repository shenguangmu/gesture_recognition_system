#!/usr/bin/env python3
"""
test_manual.py —— 校验生成出来的 Word 手册**内容完整**（不是"文件存在"）

【为什么需要它】

`build_manual.py` 跑完会打印"已生成"，但**那不说明渲染是对的**。
一个渲染器最典型的失败模式是**静默丢内容**：
某个写法没被识别，那一段就悄悄没了 ——
而这是要打印出来照着做的操作手册，**漏一行可能就是漏一个烧板前的检查**。

所以这里做的是**反向核对**：拿源 Markdown 和生成的 docx 对账，
看有没有东西在渲染过程中消失。

跑法：
    python scripts/build_manual.py      # 先出 docx
    python scripts/test_manual.py       # 再校验

判定：`*** MANUAL TESTS PASSED ***`
"""

import io
import re
import sys
from pathlib import Path

# ⚠ 只在**当前不是 UTF-8** 时才包（中文 Windows 控制台是 GBK）。
#   ⚠⚠ 那个 if 判断不能省：本文件可能被别的脚本 import，或自己 import
#   别的也会包的模块 —— 两层都包会让它们共享同一 buffer，
#   其中一层被 GC 时 buffer 被关掉，最后打印汇总时抛
#   `ValueError: I/O operation on closed file`。
#   完整说明见 host/README.md「控制台 UTF-8 兜底」一节。
if getattr(sys.stdout, 'encoding', '').lower() not in ('utf-8', 'utf8'):
    try:
        sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8',
                                      errors='replace')
    except Exception:
        pass

sys.path.insert(0, str(Path(__file__).resolve().parent))
import mdblock  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
MD = ROOT / 'docs' / 'board-bringup-guide.md'
DOCX = ROOT / 'docs' / '上板测试操作手册.docx'

_p = _f = 0


def ck(cond, msg):
    global _p, _f
    if cond:
        _p += 1
    else:
        _f += 1
        print("    [FAIL] %s" % msg)
    return cond


def _strip_probe(s):
    """把一行 Markdown 剥成"特征串"，用于在 docx 里反查。

    ⚠ 要反复剥：`> ### 标题` 同时有引用和标题。
    """
    p = s
    for _ in range(4):
        b = p
        p = re.sub(r'^>\s*', '', p)
        p = re.sub(r'^#+\s*', '', p)
        p = re.sub(r'^[-*+]\s+', '', p)
        p = re.sub(r'^\d+[.)]\s+', '', p)
        if p == b:
            break
    return p.strip('|').strip()


def main():
    print("=" * 69)
    print("  手册渲染完整性校验")
    print("=" * 69)

    if not MD.is_file():
        print("  找不到源文件：%s" % MD)
        return 2
    if not DOCX.is_file():
        print("  找不到 docx：%s" % DOCX)
        print("  先跑： python scripts/build_manual.py")
        return 2

    try:
        from docx import Document
    except ImportError:
        print("  需要 python-docx： pip install python-docx")
        return 2

    md_text = MD.read_text(encoding='utf-8')
    blocks = mdblock.parse(md_text)

    d = Document(str(DOCX))

    # 把 docx 的全部文字（段落 + 表格单元格）拼成一个大串
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    docx_text = '\n'.join(parts)

    # 源文档的块（摊平）
    flat = list(mdblock.iter_text(blocks))

    print("\n[1] 基本规模")
    ck(len(d.paragraphs) > 200, "段落数 > 200（实得 %d）" % len(d.paragraphs))
    ck(len(d.tables) >= 15, "表格数 >= 15（实得 %d）" % len(d.tables))
    # ⚠ 阈值别拍脑袋定太高：本手册正文约 1.7 万字。
    #   初版写了 2 万，直接误报失败 —— **测试的期望值也要有依据**。
    ck(len(docx_text) > 12000, "总字数 > 1.2 万（实得 %d）" % len(docx_text))

    print("\n[2] 手填记录表（打印手册的核心价值）")
    n_fill = sum(1 for p in d.paragraphs if '现场记录' in p.text)
    ck(n_fill >= 6, "至少 6 张手填记录表（实得 %d）" % n_fill)

    print("\n[3] 内容零丢失（**最重要的一条**）")
    # ⚠ 与 test_mdblock 同样的思路：源文档里所有有实质内容的行，
    #   都应该能在 docx 里找到。抓的是"整段被渲染器吞掉"。
    #
    #   ⚠ 比对前必须**去掉两边的 Markdown 记号**：
    #     源里的 `code` / **粗** / ~~删除~~ 在 docx 里是**格式**不是文字，
    #     直接比会整行误报。本校验第一版就是这么误报了 57 行 ——
    #     实际一个字都没丢（比如 `~~B1~~` 渲染成了带删除线的 "B1"）。
    #     判断这类失败时，**先去 docx 里找那一行**，别急着改渲染器。
    def _norm(s):
        # ⚠ 还要去掉**所有空白**再比。
        #   渲染器会给徽标加前后空格做内边距（'  【终端：X】  '），
        #   而 `.hwh` / `4,045,692` 这类文本本身不含空格 ——
        #   所以"去掉全部空白后相等"是安全的，且能消掉这类装饰性空格。
        return re.sub(r'\s+', '', re.sub(r'[`*~]', '', s))

    sep_re = re.compile(r'^\|[\s:|-]+\|$')
    docx_norm = _norm(docx_text)
    missing = []
    for raw in md_text.split('\n'):
        s = raw.strip()
        if not s or s == '---' or s.startswith('<!--'):
            continue

        # ⚠ 先剥块级记号（`>` / `#` / 列表符），**再**判断是不是表格行。
        #   顺序反了的话，引用块里的表格（本手册 §5.2 就有一张
        #   `> | 偏移 | 名称 | 作用 |`）会因为以 `>` 开头而被当成普通行，
        #   整行拼出来又和 docx 里的单元格对不上 —— 误报成丢失。
        #   （这是本校验踩的第 3 个坑；前两个是 ~~删除线~~ 和普通表格行。）
        body = s
        for _ in range(4):
            b2 = body
            body = re.sub(r'^>\s*', '', body)
            body = re.sub(r'^#+\s*', '', body)
            body = re.sub(r'^[-*+]\s+', '', body)
            body = re.sub(r'^\d+[.)]\s+', '', body)
            if body == b2:
                break

        if sep_re.match(body):
            continue

        # 表格行 → 逐单元格比；普通行 → 整行比
        if body.startswith('|'):
            cells = [c.strip() for c in body.strip('|').split('|')]
        else:
            cells = [body]

        for cell in cells:
            probe = cell[:24]
            if len(probe) < 12:
                continue
            if probe not in docx_text and _norm(probe) not in docx_norm:
                missing.append(s[:70])
                break

    ck(not missing,
       "源文档没有整行内容在 docx 里丢失（丢 %d 行%s）"
       % (len(missing), ('，例：' + missing[0]) if missing else ''))

    print("\n[4] 关键锚点（上板当天要照做的）")
    for key in ['引脚万用表复核', 'io_xclk', 'AP_START 在 0x00',
                'MT41K256M16', 'BRINGUP CHECK PASSED', 'gesture_system.bit',
                'bd_video.hwh', 'S2MM', 'ap_done']:
        ck(key in docx_text, "含「%s」" % key)

    print("\n[5] 每个二级标题都另起一页")
    n_h2 = sum(1 for b in blocks
               if b['type'] == 'heading' and b['level'] == 2)
    n_break = 0
    for p in d.paragraphs:
        for r in p.runs:
            if 'w:br' in r._element.xml and 'page' in r._element.xml:
                n_break += 1
    ck(n_break >= n_h2 - 1,
       "分页符 %d 个 >= 二级标题 %d 个减一" % (n_break, n_h2))

    print("\n" + "=" * 69)
    if _f == 0:
        print("  *** MANUAL TESTS PASSED ***  (%d 项)" % _p)
    else:
        print("  *** MANUAL TESTS FAILED ***  (%d 通过, %d 失败)" % (_p, _f))
    print("=" * 69)
    return 1 if _f else 0


if __name__ == '__main__':
    sys.exit(main())
